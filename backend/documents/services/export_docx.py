"""
Export CV or Letter to Word (.docx) using python-docx.
"""
import io
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build_cv_docx(cv) -> io.BytesIO:
    """Build .docx buffer for a CV model instance."""
    from .services.export_pdf import _format_experience, _format_education, _format_skills
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run(cv.full_name or 'Name')
    run.bold = True
    run.font.size = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact = ' | '.join(filter(None, [cv.email, cv.phone, cv.location]))
    if contact:
        doc.add_paragraph(contact)
    doc.add_paragraph()
    if cv.summary:
        p = doc.add_paragraph()
        p.add_run('Summary').bold = True
        doc.add_paragraph(cv.summary)
    exp_text = _format_experience(cv.experience_structured) if cv.experience_structured else (cv.experience or '')
    if exp_text:
        p = doc.add_paragraph()
        p.add_run('Experience').bold = True
        doc.add_paragraph(exp_text)
    edu_text = _format_education(cv.education_structured) if cv.education_structured else (cv.education or '')
    if edu_text:
        p = doc.add_paragraph()
        p.add_run('Education').bold = True
        doc.add_paragraph(edu_text)
    sk_text = _format_skills(cv.skills_structured) if cv.skills_structured else (cv.skills or '')
    if sk_text:
        p = doc.add_paragraph()
        p.add_run('Skills').bold = True
        doc.add_paragraph(sk_text)
    if cv.certifications:
        p = doc.add_paragraph()
        p.add_run('Certifications').bold = True
        cert_lines = [(c.get('name') or '') + (f" – {c.get('issuer', '')}" if c.get('issuer') else '') for c in cv.certifications if c.get('name')]
        doc.add_paragraph('\n'.join(cert_lines))
    if cv.projects:
        p = doc.add_paragraph()
        p.add_run('Projects').bold = True
        proj_lines = [(pr.get('name') or '') + (f" – {pr.get('description', '')}" if pr.get('description') else '') for pr in cv.projects if pr.get('name')]
        doc.add_paragraph('\n'.join(proj_lines))
    if cv.languages:
        p = doc.add_paragraph()
        p.add_run('Languages').bold = True
        lang_lines = [f"{l.get('language', '')} ({l.get('proficiency', '')})" for l in cv.languages if l.get('language')]
        doc.add_paragraph(', '.join(lang_lines))
    if cv.awards:
        p = doc.add_paragraph()
        p.add_run('Awards').bold = True
        award_lines = [(a.get('name') or '') + (f" – {a.get('issuer', '')}" if a.get('issuer') else '') for a in cv.awards if a.get('name')]
        doc.add_paragraph('\n'.join(award_lines))
    if cv.references:
        p = doc.add_paragraph()
        p.add_run('References').bold = True
        ref_lines = [f"{r.get('name', '')}, {r.get('title', '')} at {r.get('company', '')}" for r in cv.references if r.get('name')]
        doc.add_paragraph('\n'.join(ref_lines))
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_letter_docx(letter) -> io.BytesIO:
    """Build .docx buffer for a Letter model instance."""
    doc = Document()
    if letter.sender_name or letter.sender_email or letter.sender_phone:
        doc.add_paragraph(letter.sender_name or '')
        doc.add_paragraph(letter.sender_email or '')
        doc.add_paragraph(letter.sender_phone or '')
        doc.add_paragraph()
    doc.add_paragraph(f"To: {letter.recipient_name or 'Recipient'}")
    if letter.recipient_title:
        doc.add_paragraph(letter.recipient_title)
    if letter.company:
        doc.add_paragraph(letter.company)
    doc.add_paragraph()
    doc.add_paragraph(letter.body)
    doc.add_paragraph()
    doc.add_paragraph(f"Sincerely,\n{letter.sender_name or 'Sender'}")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
